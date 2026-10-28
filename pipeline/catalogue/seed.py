#!/usr/bin/env python3
"""Seed catalogue.json from what already exists:

  python3 pipeline/catalogue/seed.py ANANDPUR.xlsx

1. The Anandpur workbook (Anandpur_Library_Holdings_and_Downloads.xlsx, Drive):
   Overview → one work per row; Downloads → versions, merged where several works
   share one file (an anthology); "In my Drive" → extra links to Drive files.
2. Every PDF in the Drive inventory (../scans/inventory/inventory.json) becomes an
   in-drive version. Files no work claims are grouped by title into works marked
   needs-curation; manifest.yaml supplies titles, authors and languages where it has them.

One-off: after seeding, catalogue.json is edited directly (by hand and by sweeps).
"""

import json
import os
import unicodedata
import re
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from schema import HISTORIC_TYPES, parse_date, save, slug, validate  # noqa: E402,F401

SCANS = os.path.join(HERE, "..", "scans")

TYPE_MAP = {
    "manuscript scan": "manuscript",
    "early print (pre-1947)": "early-print",
    "critical / scholarly edition": "edition",
    "anthology containing the text": "anthology",
    "study / article": "study",
    "translation": "translation",
    "catalogue record": "catalogue-record",
    "commentary": "commentary",
}
AVAILABILITY_MAP = {
    "free download": "free-download",
    "free view only": "free-view",
    "catalogue record only (no scan)": "none",
    "login required": "login",
    "paid": "paid",
}
# Inventory file names often start with a type: "MS - …", "Early print - …".
TITLE_PREFIX_TYPES = [
    (r"ms|mss|manuscript", "manuscript"),
    (r"lithograph", "lithograph"),
    (r"early print", "early-print"),
    (r"translation", "translation"),
    (r"anthology", "anthology"),
    (r"study|article", "study"),
    (r"commentary|steek|teeka", "commentary"),
    (r"edition|critical edition", "edition"),
    (r"catalogue|catalog", "reference"),
    (r"source", None),  # the source text a translation was made from
]
KIND_TYPES = {"manuscript": "manuscript", "lithograph": "lithograph", "commentary": "commentary",
              "compilation": "compilation", "print": "edition"}
INVENTORY_SCRIPTS = {"gurmukhi": "Gurmukhi", "devanagari": "Devanagari", "arabic": "Perso-Arabic",
                     "latin": "Roman"}
# manifest.yaml language codes → (language, script implied or None)
MANIFEST_LANGS = {
    "punjabi": ("Panjabi", None), "braj": ("Braj/Bhakha", None), "english": ("English", None),
    "persian": ("Persian", None), "persian-in-gurmukhi": ("Persian", "Gurmukhi"),
    "sanskrit-in-gurmukhi": ("Sanskrit", "Gurmukhi"), "sant-bhasha": ("Sant Bhasha", None),
    "sadh-bhasha": ("Sant Bhasha", None), "sadhukari": ("Sadhukari", None),
    "sadhukari-hindi-in-gurmukhi": ("Sadhukari", "Gurmukhi"), "sanskritized-braj": ("Braj/Bhakha", None),
}


# ---------------------------------------------------------------- helpers

def split_label(label: str):
    """'Book (1976), pp. 88–91 (note)' → ('Book (1976)', 'pp. 88–91 (note)')."""
    m = re.search(r",\s*(pp?\.\s.*)$", label)
    if not m:
        return label.strip(), None
    return label[:m.start()].strip(), m.group(1).strip()


def split_list(text) -> list[str]:
    return [p.strip() for p in str(text).split(",") if p.strip()] if text else []


def drive_id_from_url(url):
    if not url:
        return None
    m = re.search(r"/d/([A-Za-z0-9_-]+)", url) or re.search(r"[?&]id=([A-Za-z0-9_-]+)", url)
    return m.group(1) if m else None


def map_type(anandpur_type: str, label: str) -> str:
    t = TYPE_MAP.get((anandpur_type or "").strip().lower(), "edition")
    if t == "early-print" and re.search(r"lithograph|litho\b|stone print", label, re.I):
        return "lithograph"
    return t


def norm_title(t: str) -> str:
    """A loose key for matching titles across romanizations."""
    t = unicodedata.normalize("NFKD", t).encode("ascii", "ignore").decode().lower()
    t = re.sub(r"\(.*?\)|&", " ", t)
    t = re.sub(r"\b(sri|shri|ji|ki|ke|di|da|de|the|of|and)\b", " ", t)
    t = re.sub(r"[^a-z0-9]+", "", t)
    return re.sub(r"([a-z])\1+", r"\1", t)  # vaaran → varan


def inventory_work_title(filename: str) -> str:
    """Best guess at the work from a Drive file name: drop '.pdf', a leading type
    ('MS - '), then everything after the first ' - ', ',' or '('."""
    t = re.sub(r"\.pdf$", "", filename, flags=re.I).strip()
    for pattern, _ in TITLE_PREFIX_TYPES:
        t = re.sub(rf"^(?:{pattern})\s+-\s+", "", t, flags=re.I)
    t = re.split(r"\s+-\s+|,|:|\(", t, maxsplit=1)[0]
    return t.strip() or filename


def type_from_filename(filename: str):
    for pattern, typ in TITLE_PREFIX_TYPES:
        if re.match(rf"^(?:{pattern})\s+-\s+", filename, re.I):
            return typ
    return None


def unique_id(base: str, taken: set) -> str:
    i, n = base, 2
    while i in taken:
        i, n = f"{base}-{n}", n + 1
    taken.add(i)
    return i


def _clean(d: dict) -> dict:
    return {k: v for k, v in d.items() if v not in (None, "", [], {})}


# ---------------------------------------------------------------- Anandpur

def from_anandpur(overview: list[dict], downloads: list[dict], found_on: str) -> dict:
    taken: set[str] = set()
    works, by_no = [], {}
    for r in overview:
        best = (r["best"] or r["printed"]).strip()
        variants = [r["printed"].strip()] if r.get("printed") and r["printed"].strip() != best else []
        notes = " ".join(x for x in (r.get("summary"), r.get("notes")) if x)
        w = _clean({
            "id": unique_id(slug(best, "w"), taken),
            "title": _clean({"roman": best, "variants": variants}),
            "authors": [{"name": a.strip()} for a in (r.get("authors") or "").split("·") if a.strip()],
            "status": "catalogued",
            "known_from": [r["source"]] if r.get("source") else [],
            "refs": {"anandpur": r["no"]},
            "notes": notes,
        })
        works.append(w)
        by_no[r["no"]] = w["id"]

    versions, by_key = [], {}
    for r in downloads:
        work_id = by_no.get(r["no"])
        if not work_id:
            continue
        label, locator = split_label(r["version"] or "")
        drive_id = drive_id_from_url(r.get("drive_url")) if r.get("saved") == "Saved" else None
        # Rows for one file merge: same Drive copy or same direct file. An item page
        # can hold several files, so it is a key only when there is no direct file.
        keys = [k for k in (drive_id, r.get("file_url")) if k] or [r.get("item_url") or label]
        entry = _clean({"work_id": work_id, "locator": locator, "note": r.get("notes")})
        v = next((by_key[k] for k in keys if k in by_key), None)
        if v is not None:
            if all(c["work_id"] != work_id for c in v["contains"]):
                v["contains"].append(entry)
            if r.get("my_notes"):
                v["user_notes"] = "; ".join(x for x in (v.get("user_notes"), r["my_notes"]) if x)
            if drive_id and not v.get("drive_id"):
                v.update(drive_id=drive_id, status="in-drive")
            for k in keys:
                by_key.setdefault(k, v)
            continue
        availability = AVAILABILITY_MAP.get((r.get("availability") or "").lower())
        if drive_id:
            status = "in-drive"
        elif availability == "none" or map_type(r.get("type"), label) == "catalogue-record":
            status = "catalogue-only"
        else:
            status = "available"
        date = parse_date(r.get("year"))
        v = _clean({
            "id": unique_id(slug(label + (f" {date['from']}" if date and date.get("from") else ""), "v"), taken),
            "label": label,
            "type": map_type(r.get("type"), label),
            "contains": [entry],
            "scripts": [s for s in split_list(r.get("scripts"))],
            "languages": split_list(r.get("languages")),
            "date": date,
            "contributors": r.get("editor"),
            "host": r.get("host"),
            "availability": availability,
            "item_url": r.get("item_url"),
            "file_url": r.get("file_url"),
            "link_check": r.get("link_check"),
            "status": status,
            "drive_id": drive_id,
            "machine_readable": r.get("machine_readable"),
            "found_on": found_on,
            "found_by": "anandpur-workbook",
            "user_notes": r.get("my_notes"),
        })
        for k in keys:
            by_key[k] = v
        versions.append(v)
    return {"version": 1, "works": works, "versions": versions}


# ---------------------------------------------------------------- Drive files

def add_drive_files(cat: dict, inventory: list[dict], manifest: list[dict], in_my_drive: list[dict]) -> None:
    """Make every inventory file an in-drive version (idempotent)."""
    taken = {w["id"] for w in cat["works"]} | {v["id"] for v in cat["versions"]}
    by_drive = {v["drive_id"]: v for v in cat["versions"] if v.get("drive_id")}
    manifest_by_id = {m["drive_id"]: m for m in manifest if m.get("drive_id")}
    by_anandpur = {w["refs"]["anandpur"]: w["id"] for w in cat["works"] if w.get("refs", {}).get("anandpur")}
    by_title = {}
    for w in cat["works"]:
        for t in [w["title"]["roman"], *w["title"].get("variants", []), w["title"].get("english", "")]:
            if t:
                by_title.setdefault(norm_title(t), w["id"])

    extra_links = defaultdict(list)  # file title → [(work_id, note)]
    for r in in_my_drive:
        if r.get("match") in ("definite", "contains") and r["no"] in by_anandpur:
            extra_links[r["file"]].append((by_anandpur[r["no"]], r.get("note")))

    for f in inventory:
        did = f["drive_id"]
        m = manifest_by_id.get(did, {})
        v = by_drive.get(did)
        if v is None:
            v = {"id": None, "label": re.sub(r"\.pdf$", "", f["title"], flags=re.I), "contains": [],
                 "status": "in-drive", "drive_id": did, "found_by": "drive-inventory"}
            v["id"] = unique_id(slug(v["label"], "v"), taken)
            v["type"] = type_from_filename(f["title"]) or KIND_TYPES.get(f.get("kind"), "edition")
            cat["versions"].append(v)
            by_drive[did] = v
        # enrich, never overwrite
        pages = m.get("pages") or f.get("pages")
        if pages and not v.get("pages"):
            v["pages"] = pages
        v.setdefault("drive_title", f["title"])
        v.setdefault("drive_folder", f.get("folder"))
        if f.get("text_layer"):
            v.setdefault("text_layer", f["text_layer"])
        langs, scripts = [], []
        for code in m.get("languages") or []:
            lang, script = MANIFEST_LANGS.get(code, (code, None))
            if lang not in langs:
                langs.append(lang)
            if script and script not in scripts:
                scripts.append(script)
        if not scripts:
            scripts = [INVENTORY_SCRIPTS[s] for s in f.get("scripts") or {} if s in INVENTORY_SCRIPTS]
            if f.get("text_layer") == "legacy" and "Gurmukhi" not in scripts:
                scripts.append("Gurmukhi")  # legacy fonts are Gurmukhi fonts
        if langs and not v.get("languages"):
            v["languages"] = langs
        if scripts and not v.get("scripts"):
            v["scripts"] = scripts
        if m.get("year") and not v.get("date"):
            d = parse_date(m["year"])
            if d and d.get("from"):
                v["date"] = d
        for k_src, k_dst in (("publisher", "publisher"), ("editor_or_commentator", "contributors")):
            val = m.get(k_src)
            if val and not str(val).lower().startswith(("unknown", "none")) and not v.get(k_dst):
                v[k_dst] = val

        linked = {c["work_id"] for c in v["contains"]}
        for work_id, note in extra_links.get(f["title"], []):
            if work_id not in linked:
                v["contains"].append(_clean({"work_id": work_id, "note": note}))
                linked.add(work_id)
        if v["contains"]:
            continue

        title_m = m.get("title") or {}
        roman = inventory_work_title(title_m.get("english") or f["title"])
        work_id = by_title.get(norm_title(roman)) or by_title.get(norm_title(inventory_work_title(f["title"])))
        if not work_id:
            work = _clean({
                "id": unique_id(slug(roman, "w"), taken),
                "title": _clean({"roman": roman, "gurmukhi": title_m.get("gurmukhi")}),
                "authors": [{"name": m["author"]}] if m.get("author") else [],
                "status": "needs-curation",
                "known_from": ["Drive"],
            })
            cat["works"].append(work)
            work_id = work["id"]
            by_title[norm_title(roman)] = work_id
        v["contains"].append({"work_id": work_id})


# ---------------------------------------------------------------- workbook

def read_anandpur(path: str):
    import openpyxl  # only the CLI needs it
    wb = openpyxl.load_workbook(path)

    def rows(name):
        ws = wb[name]
        header = [c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2):
            if row[0].value is None:
                continue
            yield header, row

    overview = []
    for h, row in rows("Overview"):
        get = {k: c.value for k, c in zip(h, row)}
        overview.append({"no": get["#"], "source": get["Source list"], "printed": get["Title (as printed)"],
                         "best": get["Best-attested title"], "authors": get["Possible authors"],
                         "summary": get["What's available (summary)"], "notes": get["Notes"]})
    downloads = []
    for h, row in rows("Downloads"):
        get = {k: c for k, c in zip(h, row)}
        val = lambda k: get[k].value if k in get else None  # noqa: E731
        link = lambda k: get[k].hyperlink.target if k in get and get[k].hyperlink else val(k)  # noqa: E731
        downloads.append({
            "no": val("#"), "version": val("Version"), "type": val("Type"), "year": val("Year"),
            "editor": val("Editor / scribe"), "host": val("Host"), "availability": val("Availability"),
            "item_url": link("Item page"), "file_url": link("Direct file"), "link_check": val("Link check"),
            "notes": val("Notes"), "saved": val("In my Drive?"), "drive_url": link("Drive copy"),
            "languages": val("Languages"), "scripts": val("Script"),
            "machine_readable": val("Machine-readable text (checked for saved files)"), "my_notes": val("My notes"),
        })
    in_my_drive = []
    for h, row in rows("In my Drive"):
        get = {k: c.value for k, c in zip(h, row)}
        in_my_drive.append({"no": get["#"], "match": get["Match"], "file": get["Drive file"], "note": get["Note"]})
    return overview, downloads, in_my_drive


def main(xlsx: str) -> None:
    import yaml
    overview, downloads, in_my_drive = read_anandpur(xlsx)
    cat = from_anandpur(overview, downloads, found_on="2026-10-08")
    with open(os.path.join(SCANS, "inventory", "inventory.json"), encoding="utf-8") as f:
        inventory = json.load(f)["files"]
    with open(os.path.join(SCANS, "manifest.yaml"), encoding="utf-8") as f:
        manifest = yaml.safe_load(f)["files"]
    add_drive_files(cat, inventory, manifest, in_my_drive)
    errors = validate(cat)
    if errors:
        print("\n".join(errors[:40]))
        sys.exit(f"{len(errors)} validation error(s); not saved")
    save(cat)
    print(f"{len(cat['works'])} works, {len(cat['versions'])} versions")


if __name__ == "__main__":
    main(sys.argv[1])
