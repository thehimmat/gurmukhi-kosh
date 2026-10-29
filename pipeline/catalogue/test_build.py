#!/usr/bin/env python3
"""Tests for the catalogue views (CATALOGUE.md and the Google Sheet export).

Run from the project root:
  python3 pipeline/catalogue/test_build.py
"""

import os
import sys
import tempfile
import unittest

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from build import composed_by, holding, render_markdown, version_rows, wanted, work_rows, write_csvs, write_xlsx  # noqa: E402


def work(i, **kw):
    return {"id": f"w-{i}", "title": {"roman": i.title()}, "status": "catalogued", **kw}


def version(i, work_ids, typ, status, **kw):
    v = {"id": f"v-{i}", "label": i, "type": typ, "status": status,
         "contains": [{"work_id": f"w-{w}"} for w in work_ids], **kw}
    if status == "in-drive":
        v.setdefault("drive_id", f"D{i}")
    return v


CAT = {
    "version": 1,
    "works": [work("gur-sobha", authors=[{"name": "Sainapati"}]), work("lav-kus"), work("pandav"),
              work("rajniti"), work("ghazals", status="needs-curation"), work("chanakya")],
    "versions": [
        version("gs-ms", ["gur-sobha"], "manuscript", "in-drive", date={"text": "1780", "from": 1780, "to": 1780}),
        version("gs-ed", ["gur-sobha"], "edition", "available", item_url="https://a.org/gs"),
        version("padam", ["lav-kus", "pandav"], "anthology", "in-drive"),
        version("pandav-ms", ["pandav"], "manuscript", "available", item_url="https://a.org/pg"),
        version("raj-cat", ["rajniti"], "catalogue-record", "catalogue-only"),
        version("ghazal-ed", ["ghazals"], "edition", "in-drive", pages=71),
    ],
}
VERSIONS_BY_WORK = {}
for v in CAT["versions"]:
    for c in v["contains"]:
        VERSIONS_BY_WORK.setdefault(c["work_id"], []).append(v)


class Holding(unittest.TestCase):
    def h(self, w):
        return holding(VERSIONS_BY_WORK.get(f"w-{w}", []))

    def test_levels(self):
        self.assertEqual(self.h("gur-sobha"), "historic")
        self.assertEqual(self.h("ghazals"), "edition")
        self.assertEqual(self.h("lav-kus"), "extracts")
        self.assertEqual(self.h("rajniti"), "known-only")
        self.assertEqual(self.h("chanakya"), "nothing")

    def test_extracts_only_but_online_ms_is_still_extracts(self):
        self.assertEqual(self.h("pandav"), "extracts")

    def test_rejected_and_candidate_versions_do_not_count(self):
        vs = [version("x", ["chanakya"], "manuscript", "candidate", found_on="2026-10-28"),
              version("y", ["chanakya"], "edition", "rejected")]
        self.assertEqual(holding(vs), "nothing")

    def test_online_only(self):
        self.assertEqual(holding([version("z", ["chanakya"], "edition", "available")]), "online-only")


class Wanted(unittest.TestCase):
    def test_works_without_a_historic_copy_are_wanted_neediest_first(self):
        ids = [w["id"] for w, _ in wanted(CAT)]
        self.assertNotIn("w-gur-sobha", ids)
        self.assertEqual(ids[:2], ["w-chanakya", "w-rajniti"])
        self.assertEqual(set(ids), {"w-chanakya", "w-rajniti", "w-lav-kus", "w-pandav"})

    def test_needs_curation_works_are_not_wanted(self):
        self.assertNotIn("w-ghazals", [w["id"] for w, _ in wanted(CAT)])

    def test_modern_work_with_an_edition_is_not_wanted(self):
        cat = {"version": 1, "works": [work("darpan", date={"text": "1962", "from": 1962, "to": 1972}),
                                       work("teeka", date={"text": "1990", "from": 1990, "to": 1990})],
               "versions": [version("darpan-ed", ["darpan"], "commentary", "in-drive")]}
        self.assertEqual([w["id"] for w, _ in wanted(cat)], ["w-teeka"])

    def test_wanted_carries_the_best_online_lead(self):
        leads = {w["id"]: lead for w, lead in wanted(CAT)}
        self.assertEqual(leads["w-pandav"]["id"], "v-pandav-ms")
        self.assertIsNone(leads["w-chanakya"])


class WantedRules(unittest.TestCase):
    def test_reference_works_are_never_wanted(self):
        cat = {"version": 1, "works": [work("suchi")],
               "versions": [version("suchi-pdf", ["suchi"], "reference", "in-drive")]}
        self.assertEqual(wanted(cat), [])

    def test_rejected_works_are_never_wanted(self):
        cat = {"version": 1, "works": [work("lesson", status="rejected")], "versions": []}
        self.assertEqual(wanted(cat), [])

    def test_twentieth_century_work_needs_only_an_edition(self):
        cat = {"version": 1, "works": [work("teeka", date={"text": "20th century", "from": 1901, "to": 2000})],
               "versions": [version("teeka-ed", ["teeka"], "edition", "in-drive")]}
        self.assertEqual(wanted(cat), [])

    def test_old_work_known_only_from_a_modern_edition_stays_wanted(self):
        # Shaheed Bilas: early 19th c., held only as a 1961 edition. A copy's date is
        # no evidence the work is modern; only the work's own date can say that.
        cat = {"version": 1, "works": [work("bilas")],
               "versions": [version("bilas-1961", ["bilas"], "edition", "in-drive",
                                    date={"text": "1961", "from": 1961, "to": 1961})]}
        self.assertEqual([w["id"] for w, _ in wanted(cat)], ["w-bilas"])


class EraPriority(unittest.TestCase):
    def cat(self):
        return {"version": 1, "works": [
            work("late", date={"text": "1880", "from": 1880, "to": 1880}),
            work("undated"),
            work("modern", date={"text": "1920", "from": 1920, "to": 1920}),
            work("early", date={"text": "1769", "from": 1769, "to": 1769}),
            work("copied-early"),
        ], "versions": [
            version("ce-ms", ["copied-early"], "catalogue-record", "catalogue-only",
                    date={"text": "1801", "from": 1801, "to": 1801}),
        ]}

    def test_composed_by_uses_work_date_else_earliest_copy(self):
        cat = self.cat()
        from build import by_work
        vbw = by_work(cat)
        works = {w["id"]: w for w in cat["works"]}
        self.assertEqual(composed_by(works["w-early"], vbw["w-early"])["from"], 1769)
        self.assertEqual(composed_by(works["w-copied-early"], vbw["w-copied-early"])["from"], 1801)
        self.assertIsNone(composed_by(works["w-undated"], []))

    def test_wanted_puts_pre_1850_first_then_undated_then_later(self):
        ids = [w["id"] for w, _ in wanted(self.cat())]
        # within an era, neediest first: "early" has nothing, "copied-early" a known MS
        self.assertEqual(ids, ["w-early", "w-copied-early", "w-undated", "w-late", "w-modern"])

    def test_rows_carry_eras(self):
        rows = {r["ID"]: r for r in work_rows(self.cat())}
        self.assertEqual((rows["w-early"]["Era"], rows["w-late"]["Era"], rows["w-undated"]["Era"]),
                         ("pre-1850", "1850-1899", "undated"))
        v = version_rows(self.cat())[0]
        self.assertEqual((v["Copy era"], v["Composed"]), ("pre-1850", ""))
        rows2 = {r["ID"]: r for r in version_rows({**self.cat(), "versions": [
            version("e-ed", ["early"], "edition", "available", date={"text": "1972", "from": 1972, "to": 1972})]})}
        self.assertEqual((rows2["v-e-ed"]["Copy era"], rows2["v-e-ed"]["Composed"]), ("1900+", "1769"))


class Rows(unittest.TestCase):
    def test_work_rows(self):
        rows = {r["ID"]: r for r in work_rows(CAT)}
        r = rows["w-gur-sobha"]
        self.assertEqual((r["Holding"], r["Versions"], r["In Drive"], r["Authors"]), ("historic", 2, 1, "Sainapati"))
        self.assertEqual(r["Earliest dated version"], 1780)

    def test_version_rows_one_per_version_with_all_works(self):
        rows = {r["ID"]: r for r in version_rows(CAT)}
        self.assertEqual(len(rows), 6)
        self.assertEqual(rows["v-padam"]["Works"], "Lav-Kus; Pandav")
        self.assertEqual(rows["v-padam"]["Drive"], "https://drive.google.com/file/d/Dpadam/view")
        self.assertEqual(rows["v-ghazal-ed"]["Pages"], 71)

    def test_version_notes_join_entry_notes(self):
        cat = {"version": 1, "works": [work("a")], "versions": [
            version("x", ["a"], "catalogue-record", "catalogue-only", notes="Written at Amritsar.")]}
        cat["versions"][0]["contains"][0]["note"] = "ff. 1-9"
        self.assertEqual(version_rows(cat)[0]["Notes"], "ff. 1-9 | Written at Amritsar.")


class Outputs(unittest.TestCase):
    def test_markdown_has_counts_and_sections(self):
        md = render_markdown(CAT)
        self.assertIn("6 works (0 rejected), 6 versions", md)
        self.assertIn("## Wanted", md)
        self.assertIn("## Needs curation", md)
        self.assertIn("| Chanakya |", md)

    def test_markdown_lists_rejected_works_apart(self):
        cat = {**CAT, "works": CAT["works"] + [work("lesson", status="rejected", notes="Rejected: teaching material")]}
        md = render_markdown(cat)
        all_works = md.split("## All works")[1].split("## Rejected")[0]
        self.assertNotIn("Lesson", all_works)
        self.assertIn("## Rejected", md)
        self.assertIn("teaching material", md.split("## Rejected")[1])

    def test_csv_views(self):
        import csv
        with tempfile.TemporaryDirectory() as d:
            write_csvs(CAT, d)
            self.assertEqual(sorted(os.listdir(d)), ["versions.csv", "wanted.csv", "works.csv"])
            with open(os.path.join(d, "works.csv"), encoding="utf-8") as f:
                rows = list(csv.DictReader(f))
            self.assertEqual(len(rows), 6)
            self.assertEqual(rows[0]["Holding"], "historic")

    def test_xlsx_tabs(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as d:
            path = os.path.join(d, "c.xlsx")
            write_xlsx(CAT, path)
            wb = openpyxl.load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Read me", "Works", "Versions", "Wanted", "Needs curation"])
            self.assertEqual(wb["Works"].max_row, 7)
            self.assertEqual(wb["Wanted"].max_row, 5)


if __name__ == "__main__":
    unittest.main()
